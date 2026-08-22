"""
exports.py — styled Excel (.xlsx) report builders for LogiTrack.

Pure formatting layer: every builder takes plain data dicts (already computed by
main.py's endpoints) and returns an openpyxl Workbook. No forecasting logic here.

Design goals: clean branded headers, banded rows, sensible number formats, frozen
header panes, auto-sized columns, and per-tab layouts — not the default ugly grid.
"""
from __future__ import annotations
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Palette ──────────────────────────────────────────────────────────────────
NAVY   = "1E293B"   # header fill
SKYD   = "0284C7"   # subtitle / brand
BAND   = "F1F5F9"   # zebra stripe
WHITE  = "FFFFFF"
TEXT   = "0F172A"
GRADE  = {"A": "16A34A", "B": "D97706", "C": "E11D48", "F": "DC2626", "—": "94A3B8"}
STATUS = {"Stockout risk": "DC2626", "Reorder due": "D97706", "Dead stock": "7C3AED",
          "Overstocked": "0891B2", "Healthy": "16A34A"}

_THIN = Side(style="thin", color="E2E8F0")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
HFONT = Font(bold=True, size=10, color=WHITE)
CFONT = Font(size=10, color=TEXT)
LEFT  = Alignment(vertical="center", horizontal="left", indent=1)
CENTER = Alignment(vertical="center", horizontal="center")

CUR = '$#,##0.00'; INT = '#,##0'; PCT = '0"%"'; PCT1 = '0.0"%"'


def _fill(hexcolor):
    return PatternFill("solid", fgColor=hexcolor)


def _title(ws, title, subtitle, ncols):
    """Branded two-row banner; returns the row where the table should start."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(1, 1, title); c.font = Font(bold=True, size=15, color=WHITE)
    c.fill = _fill(NAVY); c.alignment = LEFT
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    s = ws.cell(2, 1, subtitle); s.font = Font(size=9, color="E2E8F0")
    s.fill = _fill(SKYD); s.alignment = LEFT
    ws.row_dimensions[2].height = 16
    return 4


def _table(ws, start_row, headers, rows, formats=None, color_col=None, color_map=None, colorizer=None):
    """Write a styled table. formats: {col_index0: number_format}. color_col/color_map:
    color a column's cell by its value. colorizer(col_index0, value) -> hex|None for
    arbitrary per-cell coloring (e.g. urgency)."""
    for j, h in enumerate(headers, start=1):
        c = ws.cell(start_row, j, h); c.font = HFONT; c.fill = _fill(NAVY)
        c.alignment = LEFT; c.border = BORDER
    ws.row_dimensions[start_row].height = 20
    for i, row in enumerate(rows):
        r = start_row + 1 + i
        band = _fill(BAND) if i % 2 else _fill(WHITE)
        for j, val in enumerate(row, start=1):
            c = ws.cell(r, j, val); c.font = CFONT; c.fill = band; c.border = BORDER
            c.alignment = LEFT
            if formats and (j - 1) in formats:
                c.number_format = formats[j - 1]; c.alignment = Alignment(vertical="center", horizontal="right", indent=1)
            hexc = None
            if color_col is not None and (j - 1) == color_col and color_map:
                hexc = color_map.get(str(val))
            if colorizer:
                hexc = colorizer(j - 1, val) or hexc
            if hexc:
                c.font = Font(bold=True, size=10, color=hexc)
    ws.freeze_panes = ws.cell(start_row + 1, 1)
    return start_row + 1 + len(rows)


def _kv(ws, start_row, pairs, section=None):
    """Two-column label/value block, with an optional section header row."""
    r = start_row
    if section:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c = ws.cell(r, 1, section); c.font = Font(bold=True, size=11, color=WHITE)
        c.fill = _fill(SKYD); c.alignment = LEFT; ws.row_dimensions[r].height = 18
        r += 1
    for i, (k, v) in enumerate(pairs):
        band = _fill(BAND) if i % 2 else _fill(WHITE)
        a = ws.cell(r, 1, k); a.font = Font(size=10, color="475569"); a.fill = band; a.alignment = LEFT; a.border = BORDER
        b = ws.cell(r, 2, v if v is not None else "—"); b.font = CFONT; b.fill = band; b.alignment = LEFT; b.border = BORDER
        r += 1
    return r + 1


def _autosize(ws, skip_rows=(1, 2), max_w=58):
    widths = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is None or c.row in skip_rows:
                continue
            widths[c.column] = max(widths.get(c.column, 0), len(str(c.value)))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(w + 3, 11), max_w)


def _date(ms):
    try:
        return datetime.datetime.utcfromtimestamp(ms / 1000).strftime("%Y-%m-%d")
    except Exception:
        return ""


def _ts_date(ms):
    """Forecast timestamp (ms) -> 'Mon DD, YYYY' label, or '—' when not applicable."""
    if not ms:
        return "—"
    try:
        return datetime.datetime.utcfromtimestamp(ms / 1000).strftime("%b %d, %Y")
    except Exception:
        return "—"


def _stamp():
    return "LogiTrack · generated " + datetime.date.today().strftime("%B %d, %Y")


# ═════════════════════════════════════════════════════════════════════════════
#  PER-SKU WORKBOOK
# ═════════════════════════════════════════════════════════════════════════════
def build_sku_workbook(v: dict) -> Workbook:
    wb = Workbook(); ws = wb.active; ws.title = "Overview"
    name = v.get("skuName") or v.get("skuId")
    r = _title(ws, f"{name}", f"{v.get('skuId')} · {_stamp()}", 2)

    prot = v.get("protection") or {}
    g = v.get("orderGuardrail") or {}
    r = _kv(ws, r, [
        ("Source", v.get("filename") or ("Demo data" if v.get("mode") == "demo" else "—")),
        ("Forecast engine", f"{v.get('winningModel','')} ({v.get('route','')})"),
        ("Data eligibility", v.get("status")),
        ("Days of history", v.get("daysOfHistory")),
        ("Total units sold", v.get("totalUnitsSold")),
        ("Demand pattern", v.get("demandClass")),
        ("How it sells", v.get("demandStory") or "—"),
        *([("⚠ Forecast below recent sales", (v.get("rateCheck") or {}).get("message"))]
          if v.get("rateCheck") else []),
    ], section="Product & Model")
    dr = v.get("daysUntilReorder", -1)
    in_stock = (f"{v.get('__stock')}  (ASSUMED default — no stock column in the file; "
                f"reorder figures are estimates until real stock is entered)"
                if not v.get("stockDataAvailable", True) else v.get("__stock"))
    price_val = v.get("currentPrice") or v.get("lastPrice")
    price_txt = (f"${price_val:.2f}" if (v.get("hasPrice") and price_val is not None)
                 else "Not provided (no price column in the file)")
    r = _kv(ws, r, [
        ("In stock", in_stock),
        ("Current price", price_txt),
        ("Avg daily demand", v.get("avgDailyDemand")),
        ("Days until stockout", ("Out of stock" if v.get("alreadyOut") else v.get("daysUntilStockout")) if v.get("daysUntilStockout") is not None else "N/A (sufficient)"),
        ("Projected stockout date", _ts_date(v.get("stockoutTimestamp"))),
        ("Days until reorder", "OVERDUE" if (dr is not None and dr != -1 and dr <= 0) else (dr if dr != -1 else "N/A")),
        ("Order-by (reorder) date", _ts_date(v.get("reorderTimestamp"))),
        ("Recommended order qty", v.get("orderQty")),
        ("Order math", f"target {v.get('targetInventory','?')} (cover {v.get('coverageQty','?')} + safety {v.get('safetyStock','?')}) − stock at delivery {v.get('stockAtDelivery','?')}"),
    ], section="Reorder")
    r = _kv(ws, r, [
        ("Protection level", f"{prot.get('label','—')} ({prot.get('servicePct','—')}% service)"),
        ("Recommended level", prot.get("recommended")),
        ("Recommendation basis", {"backtest": "Per-SKU backtest cheapest tier",
                                  "economics": "Per-SKU expected cost curve",
                                  "margin": "Margin heuristic",
                                  "default": "Default (cost missing)"}.get(prot.get("source"), prot.get("source") or "—")),
        # Margin needs a price AND a cost — name whichever is actually missing.
        ("Gross margin", (f"{prot.get('marginPct')}%" if prot.get("marginPct") is not None
                          else ("Not available (no price in the data)" if prot.get("costKnownRaw")
                                else "Not available (no unit cost entered)"))),
        ("Economic detail", (
            f"Backtest: {prot.get('economics', {}).get('windows')} tests; "
            f"total ${prot.get('economics', {}).get('totalCostYr')}/yr; "
            f"holding {prot.get('economics', {}).get('holdingPct')}%/yr"
            if prot.get("source") == "backtest" and prot.get("economics") else
            f"Cost curve: margin ${prot.get('economics', {}).get('marginUnit')}/unit; "
            f"holding {prot.get('economics', {}).get('holdingPct')}%/yr"
            if prot.get("economics") else "—")),
        ("Safety buffer", f"{v.get('safetyStock','?')} units (z {v.get('zScore','?')} × σ {v.get('residualStd','?')} × √lead)"),
        ("Why this level", prot.get("reason")),
    ], section="Stockout Protection")
    if g.get("active"):
        r = _kv(ws, r, [("New-product caution", g.get("reason"))], section="Guardrail")
    cm = v.get("currentMonth") or {}
    _kv(ws, r, [
        ("Sold so far this month", cm.get("unitsSoFar")),
        ("Forecast remaining", cm.get("forecastRemaining")),
        ("Last month total", cm.get("lastMonthTotal")),
    ], section="Current Month")
    ws.column_dimensions["A"].width = 24; ws.column_dimensions["B"].width = 70

    # Daily forecast tab
    wd = wb.create_sheet("Daily Forecast")
    r = _title(wd, f"{name} — Daily Forecast & History", _stamp(), 5)
    rng = {p["x"]: p["y"] for p in (v.get("chartDataRange") or [])}
    rows = []
    for p in (v.get("chartDataHistory") or []):
        rows.append([_date(p["x"]), "Actual", p["y"], "", ""])
    for p in (v.get("chartDataFuture") or []):
        lo, hi = (rng.get(p["x"]) or ["", ""])
        rows.append([_date(p["x"]), "Forecast", p["y"], lo, hi])
    # A sub-1/day forecast formatted as a whole number is a column of zeros for a product
    # that genuinely sells every couple of days. Keep two decimals when the rate is slow.
    _slow = bool(v.get("slowSeller")) or (v.get("sparseSubtype") in ("low_volume_regular", "true_intermittent"))
    _table(wd, r, ["Date", "Type", "Units", "Low (80%)", "High (80%)"], rows,
           formats={2: ('0.00' if _slow else INT), 3: INT, 4: INT})
    _autosize(wd)

    # Monthly tab
    wm = wb.create_sheet("Monthly")
    r = _title(wm, f"{name} — Monthly Forecast", _stamp(), 6)
    _bl = bool(v.get("tooNew"))   # baseline week → forecast columns blanked
    mrows = [[c.get("monthLabel"), c.get("actualsSoFar") if c.get("isCurrent") else "",
              ("—" if _bl else c.get("forecastRemaining")), ("—" if _bl else c.get("projectedTotal")),
              ("—" if _bl else c.get("rangeLow")), ("—" if _bl else c.get("rangeHigh"))]
             for c in (v.get("monthCards") or [])]
    _table(wm, r, ["Month", "Actuals so far", "Forecast remaining", "Projected total", "Low (80%)", "High (80%)"],
           mrows, formats={1: INT, 2: INT, 3: INT, 4: INT, 5: INT})
    _autosize(wm)
    return wb


# ═════════════════════════════════════════════════════════════════════════════
#  FLEET / SUMMARY HELPERS
# ═════════════════════════════════════════════════════════════════════════════
def _fleet_row(v, folder):
    dr = v.get("daysUntilReorder", -1)
    cm = v.get("currentMonth") or {}
    prot = v.get("protection") or {}
    stock = f"{v.get('__stock')} (assumed)" if not v.get("stockDataAvailable", True) else v.get("__stock")
    margin = prot.get("marginPct") if prot.get("marginPct") is not None else "n/a (no cost)"
    basis = {"backtest": "backtest", "economics": "cost curve", "margin": "margin", "default": "default"}.get(prot.get("source"), prot.get("source") or "—")
    # A SKU still in its baseline week has no trustworthy forecast: the demand/order
    # columns would otherwise leak the pooled (and possibly volume-mismatched) figures
    # the detail screen deliberately hides. Blank them and say so in the Eligibility col.
    if v.get("tooNew"):
        D = "—"
        return [
            v.get("skuId"), v.get("skuName"), folder or "Ungrouped",
            v.get("winningModel"), f"ESTABLISHING BASELINE ({v.get('ownDays')}/{v.get('baselineDays')} days)",
            stock, D, cm.get("unitsSoFar"),
            None, None, None, None,
            D, D, D,
            prot.get("servicePct"), basis, margin,
            v.get("daysOfHistory"), v.get("totalUnitsSold"),
        ]
    # Young (provisional) SKUs forecast for real but on thin history — label them and
    # downgrade a hard "OVERDUE" to "provisional" so the sheet matches the muted UI alert.
    young = bool(v.get("young"))
    elig = f"YOUNG ({v.get('ownDays')}/{v.get('youngThreshold')} · provisional)" if young else v.get("status")
    overdue_lbl = "provisional" if young else "OVERDUE"
    return [
        v.get("skuId"), v.get("skuName"), folder or "Ungrouped",
        v.get("winningModel"), elig,
        stock, v.get("avgDailyDemand"), cm.get("unitsSoFar"),
        ("Out" if v.get("alreadyOut") else (v.get("daysUntilStockout") if v.get("daysUntilStockout") is not None else None)),
        _ts_date(v.get("stockoutTimestamp")),
        (overdue_lbl if (dr is not None and dr != -1 and dr <= 0) else (dr if dr != -1 else None)),
        _ts_date(v.get("reorderTimestamp")),
        v.get("orderQty"), v.get("safetyStock"), v.get("targetInventory"),
        prot.get("servicePct"), basis, margin,
        v.get("daysOfHistory"), v.get("totalUnitsSold"),
    ]

_FLEET_HEAD = ["SKU", "Name", "Folder", "Engine", "Eligibility", "In Stock", "Avg/Day", "Sold MTD",
               "Days→Stockout", "Stockout Date", "Days→Reorder", "Reorder By", "Order Qty",
               "Safety", "Target", "Protection %", "Protection Basis", "Margin %", "Hist (d)", "Total Sold"]
_FLEET_FMT = {5: INT, 6: '#,##0.0', 7: INT, 8: INT, 12: INT, 13: INT, 14: INT, 15: PCT, 16: PCT1, 17: INT, 18: INT}


def _urgency_color(j0, val):
    """Red for overdue, amber for due soon — on the Days→Reorder column (index 10)."""
    if j0 == 10:
        if val == "OVERDUE":
            return "DC2626"
        if isinstance(val, (int, float)) and val <= 7:
            return "D97706"
    return None


def _fleet_summary_sheet(ws, items, title):
    r = _title(ws, title, _stamp(), len(_FLEET_HEAD))
    rows = [_fleet_row(it["view"], it.get("folder")) for it in items]
    _table(ws, r, _FLEET_HEAD, rows, formats=_FLEET_FMT, colorizer=_urgency_color)
    _autosize(ws)


def _reorder_sheet(ws, items):
    r = _title(ws, "Reorder Plan — soonest first", _stamp(), 8)

    def keyf(it):
        d = it["view"].get("daysUntilReorder", 99999)
        return 99999 if d == -1 or d is None else d
    rows = []
    for it in sorted(items, key=keyf):
        v = it["view"]; dr = v.get("daysUntilReorder", -1)
        if dr == -1 or dr is None:
            continue
        rows.append([v.get("skuId"), v.get("skuName"),
                     ("provisional" if v.get("young") else "OVERDUE") if dr <= 0 else f"{dr} days",
                     _ts_date(v.get("reorderTimestamp")),
                     v.get("daysUntilStockout") if v.get("daysUntilStockout", -1) != -1 else None,
                     _ts_date(v.get("stockoutTimestamp")),
                     v.get("orderQty"), v.get("__stock")])
    if not rows:
        rows = [["—", "Nothing needs reordering right now", "", "", "", "", "", ""]]
    _table(ws, r, ["SKU", "Name", "Reorder in", "Reorder By", "Days→Stockout", "Stockout Date", "Order Qty", "In Stock"],
           rows, formats={4: INT, 6: INT, 7: INT},
           colorizer=lambda j0, val: "DC2626" if (j0 == 2 and val == "OVERDUE") else None)
    _autosize(ws)


def _monthly_sheet(ws, items):
    months = []; current = None
    for it in items:
        for c in (it["view"].get("monthCards") or []):
            ml = c.get("monthLabel")
            if ml not in months:
                months.append(ml)
            if c.get("isCurrent"):
                current = ml
    r = _title(ws, "Monthly Forecast — projected units (current month includes actuals to date)",
               _stamp(), 3 + len(months))
    head = ["SKU", "Name", "Sold MTD"] + [(f"{m} (current)" if m == current else m) for m in months]
    rows = []
    for it in items:
        v = it["view"]; cm = v.get("currentMonth") or {}
        if v.get("tooNew"):   # baseline week → no forecast to report yet
            rows.append([v.get("skuId"), v.get("skuName"), cm.get("unitsSoFar")] + ["—"] * len(months))
            continue
        m = {c.get("monthLabel"): c.get("projectedTotal") for c in (v.get("monthCards") or [])}
        rows.append([v.get("skuId"), v.get("skuName"), cm.get("unitsSoFar")] + [m.get(mo) for mo in months])
    fmt = {i: INT for i in range(2, 3 + len(months))}
    _table(ws, r, head, rows, formats=fmt)
    _autosize(ws)


def build_fleet_workbook(items: list) -> Workbook:
    wb = Workbook(); ws = wb.active; ws.title = "Fleet Summary"
    _fleet_summary_sheet(ws, items, "Fleet Summary — all SKUs")
    _reorder_sheet(wb.create_sheet("Reorder Plan"), items)
    _monthly_sheet(wb.create_sheet("Monthly Forecast"), items)
    return wb


# ═════════════════════════════════════════════════════════════════════════════
#  SUPPLIERS
# ═════════════════════════════════════════════════════════════════════════════
def _supplier_kpis(sup):
    orders = sup.get("orders") or []
    def days(a, b):
        try:
            return (datetime.date.fromisoformat(a[:10]) - datetime.date.fromisoformat(b[:10])).days
        except Exception:
            return None
    completed = [o for o in orders if o.get("receivedDate") and o.get("orderedDate")]
    lts = [d for o in completed if (d := days(o["receivedDate"], o["orderedDate"])) is not None and d >= 0]
    avg = round(sum(lts) / len(lts)) if lts else None
    withexp = [o for o in completed if o.get("expectedDate")]
    var = [d for o in withexp if (d := days(o["receivedDate"], o["expectedDate"])) is not None]
    on_time = round(sum(1 for x in var if x <= 0) / len(var) * 100) if var else None
    avgvar = round(sum(var) / len(var), 1) if var else None
    in_transit = [o for o in orders if not o.get("receivedDate")]
    return avg, on_time, avgvar, len(completed), len(in_transit)


def build_suppliers_workbook(suppliers: list) -> Workbook:
    wb = Workbook(); ws = wb.active; ws.title = "Suppliers"
    r = _title(ws, "Suppliers — reliability KPIs", _stamp(), 7)
    rows = []; hist = []
    for sup in suppliers:
        avg, on_time, avgvar, ncomp, ntrans = _supplier_kpis(sup)
        rows.append([sup.get("name"), len(sup.get("skuIds") or []), ncomp, avg, on_time, avgvar, ntrans])
        for o in (sup.get("orders") or []):
            status = "Received" if o.get("receivedDate") else "In transit"
            lt = ""
            if o.get("receivedDate") and o.get("orderedDate"):
                try:
                    lt = (datetime.date.fromisoformat(o["receivedDate"][:10]) - datetime.date.fromisoformat(o["orderedDate"][:10])).days
                except Exception:
                    lt = ""
            hist.append([sup.get("name"), o.get("skuId") or "—", o.get("qty"),
                         o.get("orderedDate"), o.get("expectedDate") or "—", o.get("receivedDate") or "—", lt, status])
    if not rows:
        rows = [["No suppliers added yet", "", "", "", "", "", ""]]
    _table(ws, r, ["Supplier", "SKUs", "Completed Orders", "Avg Lead (d)", "On-Time %", "Avg Variance (d)", "In Transit"],
           rows, formats={1: INT, 2: INT, 3: INT, 4: PCT, 5: '0.0', 6: INT})
    _autosize(ws)

    wh = wb.create_sheet("Order History")
    r = _title(wh, "Order History — all suppliers", _stamp(), 8)
    if not hist:
        hist = [["—", "—", "", "", "", "", "", "No orders logged"]]
    _table(wh, r, ["Supplier", "SKU", "Qty", "Ordered", "Expected", "Received", "Lead (d)", "Status"],
           hist, formats={2: INT, 6: INT})
    _autosize(wh)
    return wb


# ═════════════════════════════════════════════════════════════════════════════
#  SCORECARD (used by the combined workbook)
# ═════════════════════════════════════════════════════════════════════════════
def _scorecard_sheet(ws, rows):
    r = _title(ws, "Inventory Scorecard", _stamp(), 8)
    out = []
    for sc in rows:
        out.append([sc.get("skuId"), sc.get("skuName"), sc.get("status"),
                    (sc.get("daysOfCover") if sc.get("daysOfCover") is not None else "365+"),
                    (round(sc.get("sellThrough") * 100) if sc.get("sellThrough") is not None else None),
                    sc.get("returnTier"),
                    (sc.get("marginPct") if sc.get("marginPct") is not None else "n/a (no cost)"),
                    (sc.get("carryingValue") if sc.get("carryingValue") is not None else "n/a (no cost)")])
    if not out:
        out = [["—", "No SKUs", "", "", "", "", "", ""]]
    _table(ws, r, ["SKU", "Name", "Status", "Days of Cover", "Sell-through %", "Profit Grade", "Margin %", "Cash Tied Up"],
           out, formats={3: INT, 4: PCT, 6: PCT1, 7: CUR}, color_col=5, color_map=GRADE)
    _autosize(ws)


def _rows_sheet(ws, title: str, cols: list, rows: list, note: str | None = None):
    """A titled block: heading, optional note, header row, data. Used by the backtest
    sheets so 'Download everything' carries the same tables the Backtest tab shows."""
    ws.append([title]); ws["A1"].font = Font(bold=True, size=13)
    if note:
        ws.append([note]); ws.cell(ws.max_row, 1).font = Font(italic=True, size=9, color="666666")
    ws.append([])
    ws.append([c[0] for c in cols])
    hdr = ws.max_row
    for cell in ws[hdr]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334155")
    for r in rows or []:
        ws.append([c[1](r) for c in cols])
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(38, len(str(c[0])) + 4))
    ws.freeze_panes = ws.cell(hdr + 1, 1)


def _backtest_sheets(wb: Workbook, bt: dict) -> None:
    """Everything the Backtest tab shows, as sheets.

    Exported WITH its caveats — window counts, cost coverage, which levels are actually
    in use. A backtest figure that loses the context of how it was measured invites
    exactly the over-reading the tab spends so much effort preventing."""
    if not bt:
        return
    ta = bt.get("tierAnalysis") or {}
    ov = bt.get("overall") or {}
    P = bt.get("params") or {}
    ran = bt.get("ranAt")
    ws = wb.create_sheet("Backtest Summary")
    meta = [
        ("Test run", (datetime.datetime.fromtimestamp(ran).strftime("%Y-%m-%d %H:%M") if ran else "—")),
        ("Products tested", bt.get("tested")), ("Simulated reorders", bt.get("forecasts")),
        ("Failed test windows", bt.get("failedCutoffs") or 0),
        ("Lead time (days)", P.get("lead")), ("Coverage window (days)", P.get("coverage")),
        ("Holding rate %/yr", (ta.get("assumptions") or {}).get("holdingPct")),
        ("Cost basis", ta.get("costBasis")),
        ("Products costed", f"{ta.get('costedSkus')} of {ta.get('totalSkus')}"),
        ("Products missing a cost", ", ".join(ta.get("uncostedSkus") or []) or "none"),
        ("Excluded - sells at or below cost", ", ".join(ta.get("lossMakingSkus") or []) or "none"),
        ("", ""),
        ("Beats a naive forecast (MASE)", ov.get("MASE")),
        ("MASE 95% CI", " – ".join(str(x) for x in (ov.get("MASE_ci") or [])) or "—"),
        ("Average miss (WAPE %)", ov.get("WAPE%")),
        ("Runs high/low %", ov.get("bias%")),
        ("Band hit rate %", ov.get("interval_cov%")),
        ("Band hit rate 95% CI", " – ".join(str(x) for x in (ov.get("interval_cov%_ci") or [])) or "—"),
        ("Stayed in stock %", ov.get("service_achieved%")),
        ("Stayed in stock 95% CI", " – ".join(str(x) for x in (ov.get("service_achieved%_ci") or [])) or "—"),
    ]
    mp = ta.get("mixedPolicy") or {}
    if mp:
        best_uniform = min([t.get("totalCost") for t in (ta.get("tiers") or [])
                            if t.get("totalCost") is not None] or [None])
        in_use = (best_uniform is not None and mp.get("totalCost") is not None
                  and mp["totalCost"] < best_uniform)
        meta += [("", ""),
                 ("Protection policy in use",
                  "a level per product" if in_use else f"{ta.get('bestTier')}% on every product"),
                 ("Per-product mix, total $/yr", mp.get("totalCost")),
                 ("Best single level, total $/yr", best_uniform),
                 ("Mix level spread", " · ".join(f"{n} at {p}%" for p, n in
                                                 sorted((mp.get("tierCounts") or {}).items()) if n)),
                 ("Hindsight would have claimed $/yr", mp.get("inSampleTotal"))]
    _rows_sheet(ws, "Backtest — summary",
                [("Item", lambda r: r[0]), ("Value", lambda r: r[1])], meta,
                "Measured by re-running the real forecasting engines at past dates and grading them "
                "against what actually sold next.")

    acc = [("Tests", lambda r: r.get("forecasts")),
           ("Avg miss % (WAPE)", lambda r: r.get("WAPE%")),
           ("vs naive (MASE)", lambda r: r.get("MASE")),
           ("Runs high/low %", lambda r: r.get("bias%")),
           ("Band hit rate %", lambda r: r.get("interval_cov%")),
           ("Stayed in stock %", lambda r: r.get("service_achieved%")),
           ("Order size err %", lambda r: r.get("order_err%"))]
    _rows_sheet(wb.create_sheet("Backtest by Product"),
                "Backtest — per product",
                [("Product", lambda r: r.get("sku")), ("From file", lambda r: r.get("source")),
                 ("Days of history", lambda r: r.get("daysHistory")),
                 ("Engine", lambda r: r.get("engine")),
                 ("Enough windows to read?", lambda r: "NO - too few" if r.get("reportable") is False else "yes"),
                 *acc], bt.get("bySku") or [],
                "Products marked 'NO - too few' were measured on too few windows to read individually.")

    if bt.get("bySource"):
        _rows_sheet(wb.create_sheet("Backtest by File"), "Backtest — per uploaded file",
                    [("File", lambda r: r.get("source")), ("Products", lambda r: r.get("products")), *acc],
                    bt["bySource"],
                    "All products are modelled together; this splits the results by upload.")

    if ta.get("tiers"):
        rk = ta.get("ranking") or {}
        _rows_sheet(wb.create_sheet("Protection Levels"), "Protection levels compared",
                    [("Level %", lambda t: t.get("tier")),
                     ("Stayed in stock %", lambda t: t.get("achievedService")),
                     ("95% CI low", lambda t: (t.get("achievedServiceCI") or [None, None])[0]),
                     ("95% CI high", lambda t: (t.get("achievedServiceCI") or [None, None])[1]),
                     ("P(cheapest) %", lambda t: next((v for v in (
                         (rk.get("pCheapest") or {}).get(t.get("tier")),
                         (rk.get("pCheapest") or {}).get(str(t.get("tier")))) if v is not None), None)),
                     ("Buffer units", lambda t: t.get("safetyUnits")),
                     ("Missed units/yr", lambda t: t.get("unitsShortYr")),
                     ("Lost profit $/yr", lambda t: t.get("stockoutCost")),
                     ("Buffer cost $/yr", lambda t: t.get("holdingCost")),
                     ("Total $/yr", lambda t: t.get("totalCost")),
                     ("Cash in buffer $", lambda t: t.get("bufferCash"))],
                    ta["tiers"],
                    "Each row assumes ONE level for every product. See the summary sheet for which "
                    "policy is actually in use.")

    if ta.get("bySku"):
        tiers = [t.get("tier") for t in (ta.get("tiers") or [])]
        cols = [("Product", lambda r: r.get("sku")),
                ("Cost known", lambda r: "yes" if r.get("costKnown") else "no"),
                ("Sells at/below cost", lambda r: "yes" if r.get("lossMaking") else "no"),
                ("Windows", lambda r: r.get("windows")),
                ("Its own best level %", lambda r: r.get("bestTier"))]
        for t in tiers:
            cols += [(f"{t}% missed units/yr", (lambda t: lambda r: ((r.get("tiers") or {}).get(str(t)) or {}).get("unitsYr"))(t)),
                     (f"{t}% buffer units", (lambda t: lambda r: ((r.get("tiers") or {}).get(str(t)) or {}).get("safetyUnits"))(t)),
                     (f"{t}% lost profit $/yr", (lambda t: lambda r: ((r.get("tiers") or {}).get(str(t)) or {}).get("profitYr"))(t)),
                     (f"{t}% total $/yr", (lambda t: lambda r: ((r.get("tiers") or {}).get(str(t)) or {}).get("totalCostYr"))(t))]
        _rows_sheet(wb.create_sheet("Levels by Product"), "Every product at every level", cols, ta["bySku"])

    if bt.get("skipped"):
        _rows_sheet(wb.create_sheet("Backtest Not Tested"), "Products not tested",
                    [("Product", lambda r: r.get("sku")), ("Reason", lambda r: r.get("reason")),
                     ("Days of history", lambda r: r.get("days"))], bt["skipped"])


def _open_pos_sheet(ws, items: list) -> None:
    """Purchase orders on the way. Previously absent from 'everything' entirely."""
    rows = []
    for it in items or []:
        v = it.get("view") or {}
        for po in (v.get("openPOs") or []):
            rows.append({"sku": v.get("skuId"), "name": v.get("skuName"), "folder": it.get("folder"), **po})
    _rows_sheet(ws, "Open purchase orders",
                [("Product", lambda r: r.get("sku")), ("Name", lambda r: r.get("name")),
                 ("Folder", lambda r: r.get("folder")), ("Supplier", lambda r: r.get("supplier")),
                 ("Units", lambda r: r.get("units") or r.get("qty")),
                 ("Ordered", lambda r: r.get("orderedOn") or r.get("ordered")),
                 ("Expected", lambda r: r.get("expected") or r.get("eta"))],
                rows, None if rows else "No open purchase orders recorded.")


def build_all_workbook(items: list, suppliers: list, scorecard_rows: list,
                       backtest: dict | None = None) -> Workbook:
    wb = Workbook(); ws = wb.active; ws.title = "Fleet Summary"
    _fleet_summary_sheet(ws, items, "Fleet Summary — all SKUs")
    _reorder_sheet(wb.create_sheet("Reorder Plan"), items)
    _monthly_sheet(wb.create_sheet("Monthly Forecast"), items)
    _scorecard_sheet(wb.create_sheet("Scorecard"), scorecard_rows or [])
    _open_pos_sheet(wb.create_sheet("Open POs"), items)
    _backtest_sheets(wb, backtest or {})
    # suppliers (reuse the two-sheet builder's content)
    sup_wb = build_suppliers_workbook(suppliers or [])
    for src_name in ("Suppliers", "Order History"):
        src = sup_wb[src_name]; dst = wb.create_sheet(src_name)
        for row in src.iter_rows():
            for cell in row:
                d = dst.cell(cell.row, cell.column, cell.value)
                if cell.has_style:
                    d.font = cell.font.copy(); d.fill = cell.fill.copy()
                    d.border = cell.border.copy(); d.alignment = cell.alignment.copy()
                    d.number_format = cell.number_format
        for k, dim in src.column_dimensions.items():
            dst.column_dimensions[k].width = dim.width
        for k, dim in src.row_dimensions.items():
            dst.row_dimensions[k].height = dim.height
        dst.freeze_panes = src.freeze_panes
        for mc in list(src.merged_cells.ranges):
            dst.merge_cells(str(mc))
    return wb
