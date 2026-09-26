"""
Holidays: dates are known, lifts are learned, nothing is counted twice.

  1. The calendar puts every holiday on the right dates, and a day belongs to one holiday.
  2. A lift is measured against the ordinary days around it, with the weekday pattern and
     any discount taken out, and shrunk by how few years it rests on.
  3. End to end: a product's upcoming Black Friday is lifted by what its past ones showed,
     the weeks around it are not, a new product borrows its peers' lift, and nothing moves
     when no holidays are picked.
  4. A product with too little price history borrows its peers' price response, so a
     promotion still moves its forecast.
  5. The buffer grows with a deep discount and with an uncertain holiday lift.

Needs prophet installed; section 3 onward runs real fits.
"""
from __future__ import annotations
import os as _os, tempfile as _tf
_tmp = _tf.mkdtemp()
_os.environ["LOGITRACK_LIVE_PRICES_PATH"] = _os.path.join(_tmp, "live_prices.json")
_os.environ["LOGITRACK_HOLIDAYS_PATH"] = _os.path.join(_tmp, "holidays.json")

import datetime as dt
import logging
import math
import sys
import time
import warnings

warnings.filterwarnings("ignore")
logging.getLogger("cmdstanpy").disabled = True
logging.getLogger("prophet").setLevel(logging.ERROR)

import numpy as np
import pandas as pd

import holiday_calendar as H

FAILURES: list[str] = []


def check(name, cond, detail=""):
    if cond:
        print(f"  ok   {name}")
    else:
        FAILURES.append(name)
        print(f"  FAIL {name} {detail}")


D = dt.date


# ── 1. Calendar ──────────────────────────────────────────────────────────────────────
def test_calendar():
    print("holidays land on the right dates")
    w = lambda k, y: H.windows_for(k, D(y, 1, 1), D(y, 12, 31), {"selected": [k]})
    check("Thanksgiving 2025 is Nov 27", w("thanksgiving", 2025)[-1] == (D(2025, 11, 27), D(2025, 11, 27)))
    check("Black Friday weekend 2026 is Nov 27-29", (D(2026, 11, 27), D(2026, 11, 29)) in w("black_friday", 2026))
    check("Cyber Monday 2025 is Dec 1", (D(2025, 12, 1), D(2025, 12, 1)) in w("cyber_monday", 2025))
    check("Easter 2026 week ends Apr 5", (D(2026, 3, 30), D(2026, 4, 5)) in w("easter", 2026))
    check("Easter 2025 week ends Apr 20", (D(2025, 4, 14), D(2025, 4, 20)) in w("easter", 2025))
    check("Super Bowl 2026 is Feb 8", (D(2026, 2, 6), D(2026, 2, 8)) in w("super_bowl", 2026))
    check("Super Bowl 2021 was Feb 7 (first Sunday then)", (D(2021, 2, 5), D(2021, 2, 7)) in w("super_bowl", 2021))
    check("Memorial Day 2026 is May 25", (D(2026, 5, 23), D(2026, 5, 25)) in w("memorial_day", 2026))
    check("Mother's Day 2026 is May 10", (D(2026, 5, 4), D(2026, 5, 10)) in w("mothers_day", 2026))
    check("Prime Day 2025 is Jul 8-11", (D(2025, 7, 8), D(2025, 7, 11)) in w("prime_day", 2025))
    added = H.windows_for("prime_day", D(2027, 1, 1), D(2027, 12, 31),
                          {"dates": {"prime_day": [["2027-07-13", "2027-07-14"]]}})
    check("a Prime Day date added in settings is used", added == [(D(2027, 7, 13), D(2027, 7, 14))], added)

    print("a day belongs to one holiday; a holiday's main day wins")
    a = H.assign_days(D(2026, 2, 1), D(2026, 3, 20), {})
    check("Feb 14 2026 stays Valentine's inside Presidents' Day weekend", a[D(2026, 2, 14)] == "valentines", a.get(D(2026, 2, 14)))
    check("Presidents' Day itself stays Presidents' Day", a[D(2026, 2, 16)] == "presidents_day")
    check("tax refund season keeps the rest", a[D(2026, 2, 20)] == "tax_refund")
    check("every day is assigned once", len(a) == len(set(a)))

    print("an uploaded sheet moved to end yesterday still finds holidays on their real dates")
    true = pd.date_range("2024-09-01", "2026-09-20")
    moved = pd.DataFrame({"ds": true + pd.Timedelta(days=5), "y": 1.0})
    a5 = H.history_assign(moved, {}, horizon_end="2027-01-10", shift=5)
    check("last year's Christmas Day is found where the moved history put it",
          a5.get(D(2025, 12, 30)) == "christmas_day" and a5.get(D(2025, 12, 25)) != "christmas_day")
    check("and this year's is on the real calendar", a5.get(D(2026, 12, 25)) == "christmas_day")

    print("custom periods")
    c = H.validate({"custom": [{"name": "Anniversary sale", "ranges": [["2024-05-10", "2024-05-12"]], "yearly": True}]})
    k = c["custom"][0]["id"]
    got = H.windows_for(k, D(2026, 1, 1), D(2026, 12, 31), c)
    check("a yearly custom period repeats", (D(2026, 5, 10), D(2026, 5, 12)) in got, got)
    for bad, msg in [({"custom": [{"name": "", "ranges": [["2025-01-01", "2025-01-02"]]}]}, "needs a name"),
                     ({"custom": [{"name": "X", "ranges": [["2025-01-05", "2025-01-02"]]}]}, "ends before"),
                     ({"ignored": {"SKU1": ["nope"]}}, "Unknown holiday"),
                     ({"dates": {"black_friday": [["2025-01-01", "2025-01-02"]]}}, "only be added")]:
        try:
            H.validate(bad); ok = False; err = ""
        except H.SettingsError as ex:
            ok, err = msg in str(ex), str(ex)
        check(f"rejected with a message: {msg}", ok, err)
    check("every calendar period is checked automatically", len(H.selected_keys({})) == len(H.CATALOG))
    one = H.for_sku({"ignored": {"SKU1": ["prime_day"]}}, "SKU1")
    check("a period switched off for one product is skipped for that product only",
          "prime_day" not in H.selected_keys(one) and "prime_day" in H.selected_keys(H.for_sku({"ignored": {"SKU1": ["prime_day"]}}, "SKU2")))
    check("a store's own period is checked too", c["custom"][0]["id"] in H.selected_keys(c))
    check("old settings (store type, picked list) are dropped, not an error",
          H.validate({"storeType": "online", "selected": ["black_friday"]}) == {"custom": [], "dates": {}, "ignored": {}})


# ── 2. Measuring a lift ──────────────────────────────────────────────────────────────
def _series(end, days, base, lift, *, weekend=1.5, disc_pct=0.0, el=-1.5, keys=("black_friday",)):
    ds = pd.date_range(end - pd.Timedelta(days=days - 1), end)
    s = {}
    a = H.assign_days((ds[0] - pd.Timedelta(days=70)).date(), ds[-1].date(), s)
    wk = np.where(ds.weekday >= 4, weekend, 1.0)
    wk = wk / wk.mean()
    hol = np.array([a.get(x) in keys for x in ds.date])
    price = np.full(days, 50.0); listed = np.full(days, 50.0)
    price[hol] = 50.0 * (1 - disc_pct)
    lifts = np.array([(lift[a.offsets[x][0]] if isinstance(lift, (list, tuple)) else lift) if h else 1.0
                      for x, h in zip(ds.date, hol)])
    y = base * wk * lifts * (price / 50.0) ** el
    df = pd.DataFrame({"ds": ds, "y": y, "price": price, "price_listed": listed})
    return df, s, a


def test_measure():
    print("a lift is measured against the days around it, net of weekday and discount")
    end = pd.Timestamp("2026-06-30")
    df, s, a = _series(end, 900, 20.0, 2.0)
    m = H.measure(df, a)
    l, n, _ = H.own_effect(m.get("black_friday"))
    check("two past Black Fridays found", n == 2, n)
    check("the lift is about 2x, not inflated by it being a weekend", abs(math.exp(l) - 2.0) < 0.1, math.exp(l))

    df, s, a = _series(end, 900, 20.0, 2.0, disc_pct=0.3)
    m = H.measure(df, a, elasticity=-1.5, ref_price=50.0)
    l, n, _ = H.own_effect(m.get("black_friday"))
    check("a 30% Black Friday discount is taken out: still about 2x, not 3.4x",
          n == 2 and abs(math.exp(l) - 2.0) < 0.15, (n, math.exp(l)))
    m0 = H.measure(df, a)
    check("with no known price response, discounted Black Fridays are skipped, not guessed",
          (m0.get("black_friday") or {}).get("skipped") == 2 and not (m0.get("black_friday") or {}).get("logs"),
          m0.get("black_friday"))

    print("each year is kept, and the days inside the window get their own shape")
    df, s, a = _series(end, 900, 20.0, [3.0, 2.0, 1.3])
    m = H.measure(df, a)
    rec = m.get("black_friday")
    check("each past year's own number is kept", len(rec["starts"]) == 2 and len(rec["logs"]) == 2, rec["starts"])
    sh = H.shape(rec)
    check("Black Friday itself beats Saturday, Saturday beats Sunday", sh and sh[0] > sh[1] > sh[2], sh)
    flat = H.shape(H.measure(_series(end, 900, 20.0, 2.0)[0], a).get("black_friday"))
    check("an even lift stays even (no shape invented)", flat is None or max(abs(x - 1) for x in flat) < 0.08, flat)
    e = H.combine(rec, None); e["shape"] = sh
    fut = [D(2026, 11, 27), D(2026, 11, 28), D(2026, 11, 29)]
    mult, _, _ = H.future_multipliers(fut, {"black_friday": e}, H.assign_days(D(2026, 11, 1), D(2026, 12, 31), {}))
    check("and the forecast follows that shape on next year's dates", mult[0] > mult[1] > mult[2], mult)

    print("a clear lift is used as measured; an unclear one is left to safety stock; peers fill in")
    L2 = math.log(2.0)
    clear = H.combine({"logs": [L2], "expected": [60.0]}, None)
    check("one clear year on plenty of sales is used at its measured size",
          clear["applied"] and abs(clear["multiplier"] - 2.0) < 1e-9, clear)
    noisy = H.combine({"logs": [L2], "expected": [4.0]}, None)
    check("one noisy year on a handful of units isn't applied", not noisy["applied"], noisy)
    check("but its possibility goes into the safety stock", noisy["se"] > 0.2, noisy["se"])
    three = H.combine({"logs": [L2] * 3, "expected": [60.0] * 3}, None)
    check("three clear years: same measured size, less uncertain",
          abs(three["multiplier"] - 2.0) < 1e-9 and three["se"] < clear["se"], (clear["se"], three["se"]))
    quiet = H.combine({"logs": [0.03, -0.02], "expected": [60.0, 60.0]}, None)
    check("a period that doesn't move it isn't applied", not quiet["applied"] and quiet["se"] <= 0.03, quiet)
    peer = lambda x: {"black_friday": {"logs": [math.log(x)] * 2, "expected": [50.0] * 2}}
    pool = H.pool_effects([peer(1.8), peer(2.2), peer(2.0)])
    new = H.combine(None, pool.get("black_friday"))
    check("a product with no past Black Friday gets its peers' typical lift",
          new and new["applied"] and new["source"] == "peers" and 1.8 < new["multiplier"] < 2.2, new)
    mixed = H.combine({"logs": [math.log(1.2)], "expected": [5.0]}, pool.get("black_friday"))
    check("one noisy year of its own is combined with the peers, not replacing them",
          mixed and 1.5 < mixed["multiplier"] < 2.0 and mixed["source"] == "own+peers", mixed)
    solid = H.combine({"logs": [math.log(1.3)] * 3, "expected": [400.0] * 3}, pool.get("black_friday"))
    check("a product with plenty of its own clear history keeps close to its own number",
          solid and solid["multiplier"] < 1.5, solid)
    check("two peers agreeing isn't proof", not H.pool_effects([peer(2.0), peer(2.0)])["black_friday"]["clear"])
    mult, se, keys = H.future_multipliers([D(2026, 11, 26), D(2026, 11, 27)], {"black_friday": three},
                                          H.assign_days(D(2026, 11, 1), D(2026, 12, 31), {}))
    check("the lift applies on the window only", mult[0] == 1.0 and mult[1] > 1.5 and keys[1] == "black_friday", mult)

    print("a change in the everyday level right beside a window is not a holiday")
    ds = pd.date_range("2026-05-01", "2026-09-25")
    a = H.assign_days(D(2026, 3, 1), D(2026, 9, 25), {})
    lvl = np.where(ds < pd.Timestamp("2026-07-20"), 40.0, 10.0)   # a big event ends as Back to school starts
    rec = H.measure(pd.DataFrame({"ds": ds, "y": lvl}), a).get("back_to_school")
    l, n, _ = H.own_effect(rec)
    check("Back to school right after the store's level fell reads as no change, not -75%",
          n == 1 and abs(l) < 0.1, (n, round(math.exp(l), 2) if n else None))
    step_peers = [{"back_to_school": H.measure(pd.DataFrame({"ds": ds, "y": lvl * f}), a).get("back_to_school")}
                  for f in (0.5, 1.0, 1.5, 2.0)]
    check("and the whole store falling at once doesn't prove it either",
          not (H.pool_effects(step_peers).get("back_to_school") or {}).get("clear"),
          H.pool_effects(step_peers).get("back_to_school"))
    bts = np.array([a.get(x) == "back_to_school" for x in ds.date])
    rec2 = H.measure(pd.DataFrame({"ds": ds, "y": np.where(bts, 30.0, 20.0)}), a).get("back_to_school")
    check("a real Back to school lift on a steady level is still measured in full",
          abs(math.exp(H.own_effect(rec2)[0]) - 1.5) < 0.1, round(math.exp(H.own_effect(rec2)[0]), 3))
    ramp = np.linspace(20.0, 30.0, len(ds)) * np.where(bts, 1.5, 1.0)
    rec3 = H.measure(pd.DataFrame({"ds": ds, "y": ramp}), a).get("back_to_school")
    check("and on a gently rising level too",
          abs(math.exp(H.own_effect(rec3)[0]) - 1.5) < 0.12, round(math.exp(H.own_effect(rec3)[0]), 3))
    near = H.measure(pd.DataFrame({"ds": ds[ds <= pd.Timestamp("2026-08-24")],
                                   "y": np.where(bts, 30.0, 20.0)[ds <= pd.Timestamp("2026-08-24")]}), a)
    check("a window without a week of ordinary days after it isn't measured yet",
          not (near.get("back_to_school") or {}).get("logs"), near.get("back_to_school"))


# ── 3-5. End to end ──────────────────────────────────────────────────────────────────
def _catalogue(M):
    T = M.today().normalize()
    end = T - pd.Timedelta(days=1)
    rng = np.random.default_rng(3)
    s = {}
    rows = []

    def add(sku, days, base, *, lift_bf=2.0, lift_xmas=1.4, prices=None, cat="Lamps", seed=0):
        ds = pd.date_range(end - pd.Timedelta(days=days - 1), end)
        a = H.assign_days((ds[0] - pd.Timedelta(days=70)).date(), (T + pd.Timedelta(days=400)).date(), s)
        wk = np.where(ds.weekday >= 4, 1.3, 1.0)
        h = np.array([{"black_friday": lift_bf, "christmas_rush": lift_xmas}.get(a.get(x), 1.0) for x in ds.date])
        p = np.full(days, 40.0) if prices is None else prices(days)
        mu = base * wk * h * (p / 40.0) ** -1.5
        y = np.random.default_rng(seed).poisson(mu)
        for d, v, pp in zip(ds, y, p):
            rows.append({"date": d.strftime("%Y-%m-%d"), "sku": sku, "sku_name": sku, "category": cat,
                         "units_sold": int(v), "price": float(pp)})

    def two_levels(n):
        p = np.full(n, 40.0); p[n // 3: n // 3 + 90] = 32.0
        return p

    add("HOLI", 800, 12.0, prices=two_levels, seed=1)
    for i in range(3):
        add(f"PEER{i}", 800, 10.0 + 2 * i, prices=two_levels, seed=10 + i)
    add("NEWP", 75, 11.0, seed=20)                       # one price, 75 days: borrows price and holidays
    M._ingest(pd.DataFrame(rows), "hol.csv", append=False, reanchor=False, auto_backtest=False)
    return T


def _fc(M, sku, events=None, lead=14, cov=30):
    M._catalog[sku]["events"] = list(events or [])
    M._rebuild(M.today(), only=sku)
    return M.get_forecast(sku_id=sku, stock=100, lead_time_days=lead, coverage_days=cov, strategy="balanced",
                          forecast_months=12, units_on_order=0, on_order_eta_days=None, unit_cost=None,
                          fees=0.0, protection=None, tz=None, stock_source="manual", stock_counted_at=None)


def _sum(v, s, e):
    day = lambda p: pd.Timestamp(p["x"], unit="ms").strftime("%Y-%m-%d")
    return sum(p["y"] for p in v["chartDataFuture"] if s <= day(p) <= e) if v else 0.0


def test_end_to_end():
    import main as M
    T = _catalogue(M)
    t = T.date()
    bf = next(w for w in H.windows_for("black_friday", t, t + dt.timedelta(days=400), {}) if w[0] > t)
    # Same weekdays one and two weeks earlier: the ordinary baseline for those days.
    wk1 = (bf[0] - dt.timedelta(days=7), bf[1] - dt.timedelta(days=7))
    wk2 = (bf[0] - dt.timedelta(days=14), bf[1] - dt.timedelta(days=14))
    iso = lambda w: (w[0].isoformat(), w[1].isoformat())

    print("switched off for a product: nothing changes for it")
    M.set_sku_holidays("HOLI", {"ignored": [k for k in H.selected_keys({})]})
    v0 = _fc(M, "HOLI")
    check("every period is listed as switched off",
          all(r["status"] == "off" for r in (v0.get("holidays") or {}).get("effects", [])), v0.get("holidays"))
    r0 = _sum(v0, *iso(bf)) / max(_sum(v0, *iso(wk1)), 1e-9)
    check("Black Friday then looks like any other weekend", r0 < 1.3, round(r0, 2))

    print("checked automatically")
    M.set_sku_holidays("HOLI", {"ignored": []})
    v = _fc(M, "HOLI")
    rep = {r["key"]: r for r in (v.get("holidays") or {}).get("effects", [])}
    check("all 22 calendar periods were checked", len(rep) == len(H.CATALOG), len(rep))
    bfr = rep.get("black_friday") or {}
    check("Black Friday is found in its own past years and checked against its peers",
          bfr.get("status") == "applied" and bfr.get("source") == "own+peers" and bfr.get("occurrences") == 2
          and bfr.get("peers", 0) >= 2, bfr)
    check("each past year is shown", len(bfr.get("years") or []) == 2 and all("pct" in y for y in bfr["years"]), bfr.get("years"))
    check("the lift is close to the real 2x", 1.65 <= bfr.get("multiplier", 0) <= 2.3, bfr)
    fake = [k for k, r in rep.items() if r["status"] == "applied" and k not in ("black_friday", "christmas_rush")]
    check("holidays that don't move it are left out (at most 2 of 20 by chance)", len(fake) <= 2, fake)
    check("the Christmas rush is found too", (rep.get("christmas_rush") or {}).get("status") == "applied", rep.get("christmas_rush"))
    check("its day-by-day numbers are listed for next time",
          len((rep.get("black_friday") or {}).get("days") or []) in (0, 3), (rep.get("black_friday") or {}).get("days"))
    r1 = _sum(v, *iso(bf)) / max(_sum(v, *iso(wk1)), 1e-9)
    check("the upcoming Black Friday weekend is lifted", r1 > 1.45, round(r1, 2))
    near = _sum(v, *iso(wk2)) / max(_sum(v0, *iso(wk2)), 1e-9)
    check("the ordinary weeks before it are not inflated (within 12%)", 0.88 < near < 1.12, round(near, 3))
    check("the next window is named for the chart", any(w["key"] == "black_friday" for w in v["holidays"]["windows"]))
    check("holiday days are held out of the baseline", v["holidays"]["heldDays"] > 0)
    check("the details say what sales history shows",
          "Sales history shows" in ((v.get("forecastDetails") or {}).get("holidays") or {}).get("text", ""))

    print("a new product borrows its peers' Black Friday")
    vn = _fc(M, "NEWP")
    nr = {r["key"]: r for r in (vn.get("holidays") or {}).get("effects", [])}.get("black_friday") or {}
    check("its lift comes from similar products", nr.get("source") == "peers" and nr.get("active"), nr)
    check("and it says so", "similar products" in nr.get("text", ""), nr.get("text"))
    rn = _sum(vn, *iso(bf)) / max(_sum(vn, *iso(wk1)), 1e-9)
    check("and its Black Friday is lifted too", rn > 1.3, round(rn, 2))

    print("a new product borrows its peers' price response")
    pr = vn.get("priceResponse") or {}
    check("price response is borrowed from similar products", pr.get("applies") and pr.get("reason") == "borrowed", pr)
    d = lambda k: (T + pd.Timedelta(days=k)).strftime("%Y-%m-%d")
    promo = [{"type": "promotion", "date": d(5), "end_date": d(11), "discount_pct": 25.0}]
    vp = _fc(M, "NEWP", promo)
    lift = _sum(vp, d(5), d(11)) / max(_sum(vn, d(5), d(11)), 1e-9)
    check("so a 25% off promotion lifts its forecast (about 1.5x at -1.5)", 1.25 < lift < 1.8, round(lift, 2))
    check("with extra safety stock because it's borrowed",
          (vp.get("effectBuffer") or {}).get("price", 0) > 0 and vp["safetyStock"] >= vn["safetyStock"],
          (vp.get("effectBuffer"), vn["safetyStock"], vp["safetyStock"]))

    print("the buffer grows with a deep discount and an uncertain holiday")
    lead = max(14, (bf[1] - t).days + 2)
    mild = _fc(M, "HOLI", [{"type": "price_change_temporary", "date": d(3), "end_date": d(10), "new_price": 36.0}], lead=lead)
    deep = _fc(M, "HOLI", [{"type": "price_change_temporary", "date": d(3), "end_date": d(10), "new_price": 20.0}], lead=lead)
    check("a discount inside known prices adds a little", 0 < mild["effectBuffer"]["price"], mild["effectBuffer"])
    check("one far past them adds more", deep["effectBuffer"]["price"] > mild["effectBuffer"]["price"] * 1.5,
          (mild["effectBuffer"], deep["effectBuffer"]))
    check("a holiday inside the order window adds its own", mild["effectBuffer"]["holiday"] > 0, mild["effectBuffer"])
    check("and the safety stock follows", deep["safetyStock"] > mild["safetyStock"], (mild["safetyStock"], deep["safetyStock"]))

    print("settings API and report")
    g = M.get_holidays()
    check("GET lists the calendar with Prime Day's dates",
          len(g["catalog"]) == len(H.CATALOG) and any(r["key"] == "prime_day" and r.get("dates") for r in g["catalog"]))
    from fastapi import HTTPException
    try:
        M.set_sku_holidays("HOLI", {"ignored": ["not_a_holiday"]}); bad = False
    except HTTPException as ex:
        bad = ex.status_code == 400
    check("switching off an unknown period is refused with 400", bad)
    out = M.set_holidays({"custom": [{"name": "Anniversary sale", "ranges": [["2025-05-09", "2025-05-11"]], "yearly": True}]})
    check("adding a store's own period refits in the background", out["refitting"] is True)
    for _ in range(900):
        if M.get_holidays()["refit"]["status"] != "running":
            break
        time.sleep(1)
    check("and the refit finishes", M.get_holidays()["refit"]["status"] == "done", M.get_holidays()["refit"])
    check("settings survive a reload", (H.HolidaySettings().get().get("custom") or [{}])[0].get("name") == "Anniversary sale")
    v2 = M.get_forecast(sku_id="HOLI", stock=100, lead_time_days=14, coverage_days=30, strategy="balanced",
                        forecast_months=3, units_on_order=0, on_order_eta_days=None, unit_cost=None, fees=0.0,
                        protection=None, tz=None, stock_source="manual", stock_counted_at=None)
    check("the new period is checked for every product",
          any(r["key"].startswith("custom:") for r in (v2.get("holidays") or {}).get("effects", [])))
    import io as _io
    from openpyxl import load_workbook
    resp = M.holiday_report_xlsx()
    body = b"".join(resp.body_iterator) if hasattr(resp, "body_iterator") and not hasattr(resp.body_iterator, "__anext__") else None
    if body is None:
        import asyncio
        async def _read():
            return b"".join([c async for c in resp.body_iterator])
        body = asyncio.run(_read())
    wb = load_workbook(_io.BytesIO(body))
    rows = list(wb["By product"].iter_rows(values_only=True))
    check("the report has a row per product per period",
          len(rows) - 1 == len(M._cache()) * (len(H.CATALOG) + 1), (len(rows), len(M._cache())))
    check("and a summary by holiday", wb["By holiday"].max_row > 1)


def test_backwards_price_and_resync():
    """A product whose own prices read backwards borrows its peers' response, and a live
    store's nightly re-sync keeps each product's planned events and hand-set category."""
    import main as M
    T = M.today().normalize()
    end = T - pd.Timedelta(days=1)
    rows = []

    def add(sku, days, base_fn, price_fn, seed, cat="Shirts"):
        ds = pd.date_range(end - pd.Timedelta(days=days - 1), end)
        p = price_fn(days)
        y = np.random.default_rng(seed).poisson(base_fn(days) * (p / 40.0) ** -1.5)
        for d, v, pp in zip(ds, y, p):
            rows.append({"date": d.strftime("%Y-%m-%d"), "sku": sku, "sku_name": sku, "category": cat,
                         "units_sold": int(v), "price": float(pp)})

    def two_levels(n):
        p = np.full(n, 40.0); p[n // 3: n // 3 + 60] = 32.0
        return p
    for i in range(3):
        add(f"SHIRT{i}", 200, lambda n, i=i: np.full(n, 8.0 + i), two_levels, 40 + i)
    # Sold at $30 through a slow spell, then at $40 once a busy season began: dearer and busier.
    add("BACKWARDS", 200, lambda n: np.where(np.arange(n) < 100, 1.0, 6.0),
        lambda n: np.where(np.arange(n) < 100, 30.0, 40.0), 50)
    # The only product in its category, and it has only ever sold at one price.
    add("GADGET", 200, lambda n: np.full(n, 5.0), lambda n: np.full(n, 40.0), 60, cat="Gadgets")
    M._ingest(pd.DataFrame(rows), "store", append=False, reanchor=False, auto_backtest=False)

    print("a product whose own prices read backwards borrows its peers' price response")
    cut = [{"type": "price_change_temporary", "date": (T + pd.Timedelta(days=2)).strftime("%Y-%m-%d"),
            "end_date": (T + pd.Timedelta(days=8)).strftime("%Y-%m-%d"), "new_price": 30.0}]
    base = _fc(M, "BACKWARDS")
    pr = base.get("priceResponse") or {}
    check("its own prices clear the gate but aren't used", pr.get("reason") == "borrowed", pr)
    check("and the page says why, in plain words", "sold more at its higher price" in (pr.get("text") or ""), pr.get("text"))
    lifted = _fc(M, "BACKWARDS", cut)
    s0 = _sum(base, (T + pd.Timedelta(days=2)).strftime("%Y-%m-%d"), (T + pd.Timedelta(days=8)).strftime("%Y-%m-%d"))
    s1 = _sum(lifted, (T + pd.Timedelta(days=2)).strftime("%Y-%m-%d"), (T + pd.Timedelta(days=8)).strftime("%Y-%m-%d"))
    check("so a 25% price cut now lifts its forecast", s1 > s0 * 1.2, (s0, s1))

    print("with no similar products to borrow from, the whole store's price response stands in")
    g = _fc(M, "GADGET")
    gp = g.get("priceResponse") or {}
    check("its price response is borrowed from the store", gp.get("reason") == "borrowed" and gp.get("scope") == "store", gp)
    check("and the page says where it came from", "other products in the store" in (gp.get("text") or ""), gp.get("text"))
    g1 = _fc(M, "GADGET", cut)
    check("so a price cut lifts it too",
          _sum(g1, (T + pd.Timedelta(days=2)).strftime("%Y-%m-%d"), (T + pd.Timedelta(days=8)).strftime("%Y-%m-%d"))
          > _sum(g, (T + pd.Timedelta(days=2)).strftime("%Y-%m-%d"), (T + pd.Timedelta(days=8)).strftime("%Y-%m-%d")) * 1.2)

    print("the nightly re-sync keeps what was set on each product")
    M._catalog["SHIRT0"]["events"] = list(cut)
    M._set_attributes_locked({"BACKWARDS": {"category": "Jerseys"}}, 0)
    M._ingest(pd.DataFrame(rows), "store", append=False, reanchor=False, auto_backtest=False)
    check("a planned price change survives the re-sync", M._catalog["SHIRT0"]["events"] == cut,
          M._catalog["SHIRT0"]["events"])
    check("so does a category set by hand", (M._catalog["BACKWARDS"].get("attrs") or {}).get("category") == "Jerseys",
          M._catalog["BACKWARDS"].get("attrs"))
    M._ingest(pd.DataFrame(rows), "sheet.xlsx", append=False, reanchor=True, auto_backtest=False)
    check("a spreadsheet upload still starts clean", not M._catalog["SHIRT0"].get("events"),
          M._catalog["SHIRT0"].get("events"))


if __name__ == "__main__":
    print("\nHolidays\n" + "-" * 44)
    for fn in (test_calendar, test_measure, test_end_to_end, test_backwards_price_and_resync):
        print(f"\n{fn.__name__}")
        fn()
    print()
    if FAILURES:
        print(f"{len(FAILURES)} FAILED: {', '.join(FAILURES)}")
        sys.exit(1)
    print("All holiday tests passed.")
